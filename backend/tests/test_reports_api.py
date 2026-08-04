"""Tests report router — sinh Word/Excel từ verdict HSDT (offline, seed DB trực tiếp).

Không gọi /rubric hay /evaluate thật (cần proxy); seed HsdtCriterionEval + HsdtVerdict qua DB.
"""
import io

import pytest
from docx import Document


@pytest.fixture
def db_session(client):  # noqa: ARG001 — phụ thuộc client để DB được khởi tạo trước
    """Session DB dùng chung engine với app."""
    import database as _db
    sess = _db.SessionLocal()
    try:
        yield sess
    finally:
        sess.close()


def _seed_verdict(db, package_id: int, vendor_id: int, ket_qua: str = "đạt") -> int:
    """Tạo 1 HsdtCriterionEval + 1 verdict; trả verdict_id."""
    import models
    ev = models.HsdtCriterionEval(
        package_id=package_id, vendor_id=vendor_id, thu_tu=0, nhom="hop_le",
        ten="Đơn dự thầu hợp lệ", ket_qua=ket_qua)
    ev.verdicts.append(models.HsdtVerdict(
        thu_tu=0, noi_dung_kiem_tra="Chữ ký & con dấu", hsdt_kiem_tra="don_du_thau",
        yeu_cau="có chữ ký", ket_qua=ket_qua, bang_chung="Có chữ ký, đóng dấu",
        trang=[1], do_tin=0.9))
    db.add(ev)
    db.commit()
    db.refresh(ev)
    return ev.verdicts[0].id


def _package(client) -> tuple[int, int]:
    p = client.post("/api/v1/packages",
                    json={"ma_so": "G-R", "ten": "Gói R", "vendors": ["NhaThauA"]}).json()["data"]
    return p["id"], p["vendors"][0]["id"]


def test_generate_and_download_word(client, db_session):
    pid, vid = _package(client)
    _seed_verdict(db_session, pid, vid)
    gen = client.post(f"/api/v1/packages/{pid}/reports?loai=word").json()["data"]
    assert gen["report_id"]
    dl = client.get(f"/api/v1/reports/{gen['report_id']}/download")
    assert dl.status_code == 200 and len(dl.content) > 0

    doc = Document(io.BytesIO(dl.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "NhaThauA" in text and "Đơn dự thầu hợp lệ" in text


def test_word_shows_hinh_thuc_nguon_and_phat_hien(client, db_session):
    """Báo cáo phản ánh cấu trúc mới: hình thức nhà thầu, điều khoản nguồn, phát hiện ngoài checklist."""
    import models
    pid, vid = _package(client)
    # tiêu chí hợp lệ có điều khoản nguồn
    ev = models.HsdtCriterionEval(package_id=pid, vendor_id=vid, thu_tu=0, nhom="hop_le",
                                  ten="Bảo đảm dự thầu", ket_qua="đạt",
                                  yeu_cau_goc="Nộp bảo đảm 6.100.000")
    ev.verdicts.append(models.HsdtVerdict(
        thu_tu=0, noi_dung_kiem_tra="Giá trị", hsdt_kiem_tra="bao_dam_du_thau", ket_qua="đạt",
        bang_chung="6.1tr", trang=[1], do_tin=0.9, nguon_hsmt="E-BDL 18.1"))
    # phát hiện bổ sung (nhóm synthetic) — KHÔNG được coi là tiêu chí hợp lệ
    ph = models.HsdtCriterionEval(package_id=pid, vendor_id=vid, thu_tu=0,
                                  nhom="phat_hien_bo_sung", ten="Phát hiện của hệ thống", ket_qua="đạt")
    ph.verdicts.append(models.HsdtVerdict(
        thu_tu=0, noi_dung_kiem_tra="Người ký khớp ĐKKD", hsdt_kiem_tra="don_du_thau",
        ket_qua="đạt", bang_chung="khớp", trang=[1], do_tin=0.9))
    db_session.add_all([ev, ph, models.HsdtVendorEval(
        package_id=pid, vendor_id=vid, hinh_thuc="độc lập", nguon="khai báo",
        bang_chung="dự thầu độc lập", do_tin=0.9)])
    db_session.commit()

    gen = client.post(f"/api/v1/packages/{pid}/reports?loai=word").json()["data"]
    dl = client.get(f"/api/v1/reports/{gen['report_id']}/download")
    doc = Document(io.BytesIO(dl.content))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "độc lập" in text                       # hình thức nhà thầu
    assert "E-BDL 18.1" in text                     # điều khoản nguồn HSMT
    assert "ngoài checklist" in text and "Người ký khớp ĐKKD" in text   # phát hiện tách riêng


def test_generate_and_download_excel(client, db_session):
    pid, vid = _package(client)
    _seed_verdict(db_session, pid, vid)
    gen = client.post(f"/api/v1/packages/{pid}/reports?loai=excel").json()["data"]
    assert gen["report_id"] and gen["loai"] == "excel"
    dl = client.get(f"/api/v1/reports/{gen['report_id']}/download")
    assert dl.status_code == 200 and len(dl.content) > 0


def test_generate_report_missing_package(client):
    r = client.post("/api/v1/packages/99999/reports?loai=word")
    assert r.status_code == 404


def test_download_missing_report(client):
    r = client.get("/api/v1/reports/99999/download")
    assert r.status_code == 404


def test_export_blocked_when_unresolved_error(client, db_session):
    """Xuất báo cáo phải trả 409 khi còn verdict ket_qua='lỗi' chưa override."""
    pid, vid = _package(client)
    _seed_verdict(db_session, pid, vid, ket_qua="lỗi")
    r = client.post(f"/api/v1/packages/{pid}/reports?loai=excel")
    assert r.status_code == 409
    assert "ai lỗi" in r.json()["error"].lower()


def test_passed_legality_chi_dat_va_khong_ap_dung(client, db_session):
    """Hợp lệ CHỈ khi mọi tiêu chí 'đạt' hoặc 'không áp dụng'.

    'cần làm rõ'/'thiếu hồ sơ'/'lỗi' = CHƯA kết luận được -> không được coi là hợp lệ.
    """
    from routers.reports import _rebuild_evals
    import models

    pid, vid = _package(client)
    pkg = db_session.get(models.ProcurementPackage, pid)

    def _passed(*ket_quas: str) -> bool:
        for e in db_session.query(models.HsdtCriterionEval).filter_by(package_id=pid).all():
            db_session.delete(e)
        db_session.commit()
        for kq in ket_quas:
            _seed_verdict(db_session, pid, vid, kq)
        db_session.expire_all()
        return _rebuild_evals(pkg, db_session)[1][vid]["passed_legality"]

    assert _passed("đạt") is True
    assert _passed("đạt", "không áp dụng") is True
    assert _passed("không áp dụng") is True
    assert _passed("đạt", "cần làm rõ") is False
    assert _passed("cần làm rõ") is False
    assert _passed("thiếu hồ sơ") is False
    assert _passed("lỗi") is False
    assert _passed("đạt", "không đạt") is False


def test_tieu_chi_thieu_ho_so_khong_duoc_tinh_la_hop_le():
    """'thiếu hồ sơ' nay là kết luận cấp tiêu chí — báo cáo KHÔNG được coi nó là hợp lệ.

    `_KET_QUA_HOP_LE` chỉ gồm {đạt, không áp dụng} nên hành vi vốn đã đúng; test này khoá lại để
    ai nới tập hợp lệ sau này phải thấy đỏ.
    """
    from routers.reports import _KET_QUA_HOP_LE

    assert "thiếu hồ sơ" not in _KET_QUA_HOP_LE
    assert "cần làm rõ" not in _KET_QUA_HOP_LE
    assert "lỗi" not in _KET_QUA_HOP_LE
